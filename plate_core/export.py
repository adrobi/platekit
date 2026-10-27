"""Export plate maps, Master Mix tables and DNA normalization to Excel (.xlsx) and printable HTML Bench Sheet."""

from __future__ import annotations

from datetime import datetime
import html
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from i18n import tr
from plate_core.mastermix import MixRecipe, calculate_plate_mixes
from plate_core.models import ROLE_COLORS, Plate, row_label
from plate_core.normalization import NormalizationConfig, calculate_normalization


def _lighten_hex(hex_color: str, factor: float = 0.78) -> str:
    """Blend a hex color (#RRGGBB) with white for readable Excel/HTML backgrounds."""
    clean = hex_color.lstrip("#")
    if len(clean) != 6:
        return "E5E7EB"
    r = int(clean[0:2], 16)
    g = int(clean[2:4], 16)
    b = int(clean[4:6], 16)
    lr = int(r + (255 - r) * factor)
    lg = int(g + (255 - g) * factor)
    lb = int(b + (255 - b) * factor)
    return f"{lr:02X}{lg:02X}{lb:02X}"


def _darken_hex(hex_color: str, factor: float = 0.68) -> str:
    """Blend a hex color (#RRGGBB) with darkzinc (#18181b) for dark-mode HTML previews."""
    clean = hex_color.lstrip("#")
    if len(clean) != 6:
        return "27272A"
    r = int(clean[0:2], 16)
    g = int(clean[2:4], 16)
    b = int(clean[4:6], 16)
    dr = int(r * (1.0 - factor) + 24 * factor)
    dg = int(g * (1.0 - factor) + 24 * factor)
    db = int(b * (1.0 - factor) + 27 * factor)
    return f"{dr:02X}{dg:02X}{db:02X}"


def export_excel_workbook(
    plate: Plate,
    recipe: MixRecipe,
    norm_config: NormalizationConfig,
    path: str | Path,
) -> None:
    """Create a multi-sheet Excel (.xlsx) report for the plate."""
    wb = openpyxl.Workbook()

    thin_side = Side(border_style="thin", color="CBD5E1")
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(bold=True, color="FFFFFF")
    subhdr_fill = PatternFill("solid", fgColor="E2E8F0")
    subhdr_font = Font(bold=True, color="0F172A")

    # --- Sheet 1: Plate Map ---
    ws_map = wb.active
    ws_map.title = "Plate Map"
    ws_map["A1"] = f"Plate: {plate.name} ({plate.format.title})"
    ws_map["A1"].font = Font(size=13, bold=True)
    ws_map["A2"] = f"Operator: {plate.operator or '—'} | Recipe: {recipe.name}"

    start_row = 4
    ws_map.cell(row=start_row, column=1, value="Row \\ Col").fill = subhdr_fill
    ws_map.cell(row=start_row, column=1).font = subhdr_font
    for c in range(plate.cols):
        cell = ws_map.cell(row=start_row, column=c + 2, value=c + 1)
        cell.fill = subhdr_fill
        cell.font = subhdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = grid_border
        ws_map.column_dimensions[get_column_letter(c + 2)].width = 15

    for r in range(plate.rows):
        row_idx = start_row + 1 + r
        ws_map.row_dimensions[row_idx].height = 36
        hdr_cell = ws_map.cell(row=row_idx, column=1, value=row_label(r))
        hdr_cell.fill = subhdr_fill
        hdr_cell.font = subhdr_font
        hdr_cell.alignment = Alignment(horizontal="center", vertical="center")
        hdr_cell.border = grid_border

        for c in range(plate.cols):
            w = plate.get_well(r, c)
            cell = ws_map.cell(row=row_idx, column=c + 2)
            cell.border = grid_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if w.is_occupied:
                lines = [w.sample or w.role.upper()]
                if w.assay:
                    lines.append(f"[{w.assay}]")
                if w.role != "unknown":
                    lines.append(f"({w.role.upper()})")
                cell.value = "\n".join(lines)
                base_hex = (
                    plate.ensure_assay_color(w.assay)
                    if w.assay
                    else ROLE_COLORS.get(w.role, "#6B7280")
                )
                cell.fill = PatternFill("solid", fgColor=_lighten_hex(base_hex, 0.72))
                cell.font = Font(size=9, bold=(w.role in ("ntc", "pos", "std")))
            else:
                cell.value = ""

    # --- Sheet 2: Master Mix ---
    ws_mix = wb.create_sheet("Master Mix")
    ws_mix["A1"] = f"Master Mix Calculation — {recipe.name}"
    ws_mix["A1"].font = Font(size=13, bold=True)
    ws_mix["A2"] = (
        f"Reaction Volume: {recipe.reaction_volume_ul:g} µL | "
        f"Template DNA: {recipe.template_volume_ul:g} µL | "
        f"Overfill: {recipe.overfill_mode} ({recipe.overfill_value:g})"
    )

    mixes = calculate_plate_mixes(plate, recipe, water_label=tr("mix.water_auto"))
    cur_row = 4
    for calc in mixes:
        ws_mix.cell(
            row=cur_row,
            column=1,
            value=(
                f"Assay: {calc.assay}  |  Wells: {calc.well_count}  |  "
                f"Effective Rxns: {calc.effective_rxns:.2f}×  |  "
                f"Dispense: {calc.dispense_per_well_ul:.2f} µL/well"
            ),
        ).font = subhdr_font
        cur_row += 1
        headers = ["Check", "Component", "Stock", "Final", "1× Vol (µL)", "Total Vol (µL)"]
        for col_i, h in enumerate(headers, start=1):
            hc = ws_mix.cell(row=cur_row, column=col_i, value=h)
            hc.fill = header_fill
            hc.font = header_font
            hc.border = grid_border
        cur_row += 1
        for line in calc.lines:
            vals = ["[  ]", line.component, line.stock_conc, line.final_conc, line.vol_1x_ul, line.vol_total_ul]
            for col_i, v in enumerate(vals, start=1):
                c_cell = ws_mix.cell(row=cur_row, column=col_i, value=v)
                c_cell.border = grid_border
            cur_row += 1
        tot_cell = ws_mix.cell(row=cur_row, column=5, value="Total Mix (µL):")
        tot_cell.font = subhdr_font
        val_cell = ws_mix.cell(row=cur_row, column=6, value=calc.total_mix_volume_ul)
        val_cell.font = subhdr_font
        cur_row += 2

    for col_i, width in enumerate((8, 30, 12, 12, 14, 16), start=1):
        ws_mix.column_dimensions[get_column_letter(col_i)].width = width

    # --- Sheet 3: DNA Normalization ---
    ws_norm = wb.create_sheet("DNA Normalization")
    ws_norm["A1"] = f"DNA Normalization ({norm_config.mode})"
    ws_norm["A1"].font = Font(size=13, bold=True)
    norm_headers = [
        "Sample",
        "Wells",
        "Stock (ng/µL)",
        "Pre-dilution",
        "DNA Vol (µL)",
        "Water/Buffer (µL)",
        "Target Result",
        "QC Status",
    ]
    for col_i, h in enumerate(norm_headers, start=1):
        hc = ws_norm.cell(row=3, column=col_i, value=h)
        hc.fill = header_fill
        hc.font = header_font
        hc.border = grid_border

    norm_rows = calculate_normalization(plate, norm_config)
    qc_fills = {
        "ok": PatternFill("solid", fgColor="DCFCE7"),
        "too_concentrated": PatternFill("solid", fgColor="FEF3C7"),
        "too_dilute": PatternFill("solid", fgColor="FEE2E2"),
        "missing_conc": PatternFill("solid", fgColor="F1F5F9"),
    }
    for idx, nrow in enumerate(norm_rows, start=4):
        predil = f"1:{nrow.predilution_factor}" if nrow.predilution_factor > 1 else "—"
        vals = [
            nrow.sample,
            ", ".join(nrow.wells),
            nrow.stock_conc_ng_ul if nrow.stock_conc_ng_ul is not None else "—",
            predil,
            nrow.dna_volume_ul,
            nrow.diluent_volume_ul,
            nrow.final_summary,
            nrow.qc_status.upper(),
        ]
        row_fill = qc_fills.get(nrow.qc_status)
        for col_i, v in enumerate(vals, start=1):
            c_cell = ws_norm.cell(row=idx, column=col_i, value=v)
            c_cell.border = grid_border
            if row_fill and col_i == 8:
                c_cell.fill = row_fill

    for col_i, width in enumerate((18, 22, 15, 14, 14, 18, 26, 20), start=1):
        ws_norm.column_dimensions[get_column_letter(col_i)].width = width

    # --- Sheet 4: Well List ---
    ws_list = wb.create_sheet("Well List")
    list_headers = [
        "Well",
        "Well (Padded)",
        "Row",
        "Col",
        "Sample",
        "DNA Conc (ng/µL)",
        "Assay",
        "Role",
        "Replicate Group",
        "Std Quantity",
        "Note",
    ]
    for col_i, h in enumerate(list_headers, start=1):
        hc = ws_list.cell(row=1, column=col_i, value=h)
        hc.fill = header_fill
        hc.font = header_font
        hc.border = grid_border

    for idx, w in enumerate(plate.occupied_wells(direction="row"), start=2):
        vals = [
            w.coord,
            w.coord_padded,
            row_label(w.row),
            w.col + 1,
            w.sample,
            w.concentration,
            w.assay,
            w.role,
            w.replicate_group,
            w.std_quantity,
            w.note,
        ]
        for col_i, v in enumerate(vals, start=1):
            c_cell = ws_list.cell(row=idx, column=col_i, value=v)
            c_cell.border = grid_border

    wb.save(Path(path))


def render_bench_sheet_html(
    plate: Plate,
    recipe: MixRecipe,
    norm_config: NormalizationConfig,
    *,
    dark: bool = False,
) -> str:
    """Render a self-contained printable A4 HTML Bench Sheet (supports dark mode for in-app preview)."""
    esc = html.escape
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    mixes = calculate_plate_mixes(plate, recipe, water_label=tr("mix.water_auto"))
    norm_rows = calculate_normalization(plate, norm_config)

    if dark:
        bg_body = "#18181b"
        fg_body = "#e4e4e7"
        fg_muted = "#a1a1aa"
        border_col = "#3f3f46"
        th_grid_bg = "#27272a"
        th_grid_fg = "#f4f4f5"
        empty_bg = "#1f1f23"
        empty_fg = "#71717a"
        coord_fg = "#a1a1aa"
        assay_fg = "#d4d4d8"
        th_data_bg = "#27272a"
        th_data_fg = "#f4f4f5"
        td_data_bg = "#1f1f23"
        td_data_fg = "#e4e4e7"
        qc_warn_bg = "#713f12"
        qc_warn_fg = "#fef3c7"
        qc_err_bg = "#7f1d1d"
        qc_err_fg = "#fee2e2"
        qc_ok_bg = "#14532d"
        qc_ok_fg = "#dcfce7"
    else:
        bg_body = "#ffffff"
        fg_body = "#0f172a"
        fg_muted = "#475569"
        border_col = "#cbd5e1"
        th_grid_bg = "#f1f5f9"
        th_grid_fg = "#0f172a"
        empty_bg = "#fafafa"
        empty_fg = "#94a3b8"
        coord_fg = "#64748b"
        assay_fg = "#334155"
        th_data_bg = "#1e293b"
        th_data_fg = "#ffffff"
        td_data_bg = "#ffffff"
        td_data_fg = "#0f172a"
        qc_warn_bg = "#fef3c7"
        qc_warn_fg = "#713f12"
        qc_err_bg = "#fee2e2"
        qc_err_fg = "#7f1d1d"
        qc_ok_bg = "#dcfce7"
        qc_ok_fg = "#14532d"

    # Plate grid HTML
    grid_rows_html: list[str] = []
    col_headers = "".join(f"<th>{c + 1}</th>" for c in range(plate.cols))
    grid_rows_html.append(f"<tr><th></th>{col_headers}</tr>")

    for r in range(plate.rows):
        cells_html: list[str] = [f"<th class='row-hdr'>{row_label(r)}</th>"]
        for c in range(plate.cols):
            w = plate.get_well(r, c)
            if w.is_occupied:
                base_hex = (
                    plate.ensure_assay_color(w.assay)
                    if w.assay
                    else ROLE_COLORS.get(w.role, "#6B7280")
                )
                bg = f"#{_darken_hex(base_hex, 0.65)}" if dark else f"#{_lighten_hex(base_hex, 0.76)}"
                title = esc(w.sample or w.role.upper())
                sub = esc(w.assay) if w.assay else ""
                role_tag = (
                    f"<span class='role-tag'>{esc(w.role.upper())}</span>"
                    if w.role != "unknown"
                    else ""
                )
                cells_html.append(
                    f"<td class='well occupied' style='background-color:{bg};color:{fg_body};border-left:3px solid {base_hex};'>"
                    f"<div class='w-coord'>{w.coord}</div>"
                    f"<div class='w-sample'>{title}</div>"
                    f"<div class='w-assay'>{sub} {role_tag}</div>"
                    f"</td>"
                )
            else:
                cells_html.append(
                    f"<td class='well empty'><div class='w-coord'>{w.coord}</div></td>"
                )
        grid_rows_html.append("<tr>" + "".join(cells_html) + "</tr>")

    # Master mix tables HTML
    col_done = esc(tr("bench.col.done"))
    col_comp = esc(tr("mix.col.component"))
    col_stock = esc(tr("mix.col.stock"))
    col_final = esc(tr("mix.col.final"))
    col_v1 = esc(tr("mix.col.vol_1x"))
    col_vtot = esc(tr("mix.col.vol_total"))
    lbl_total_mix = esc(tr("bench.total_mix"))

    mix_blocks: list[str] = []
    for calc in mixes:
        rows_str = "".join(
            f"<tr>"
            f"<td class='chk'>&#9744;</td>"
            f"<td>{esc(ln.component)}</td>"
            f"<td>{esc(ln.stock_conc)}</td>"
            f"<td>{esc(ln.final_conc)}</td>"
            f"<td class='num'>{ln.vol_1x_ul:.2f}</td>"
            f"<td class='num'><b>{ln.vol_total_ul:.2f}</b></td>"
            f"</tr>"
            for ln in calc.lines
        )
        hdr_summary = esc(
            tr(
                "mix.summary_header",
                assay=calc.assay,
                wells=calc.well_count,
                eff=calc.effective_rxns,
                total=calc.total_mix_volume_ul,
                disp=calc.dispense_per_well_ul,
                tmpl=calc.template_volume_ul,
            )
        )
        mix_blocks.append(
            f"<div class='mix-card'>"
            f"<h3>{hdr_summary}</h3>"
            f"<table class='data-table'>"
            f"<thead><tr><th>{col_done}</th><th>{col_comp}</th><th>{col_stock}</th><th>{col_final}</th><th>{col_v1}</th><th>{col_vtot}</th></tr></thead>"
            f"<tbody>{rows_str}</tbody>"
            f"<tfoot><tr><td colspan='5' style='text-align:right'><b>{lbl_total_mix}</b></td>"
            f"<td class='num'><b>{calc.total_mix_volume_ul:.2f} &micro;L</b></td></tr></tfoot>"
            f"</table></div>"
        )

    # DNA Normalization table HTML
    norm_tr: list[str] = []
    for nrow in norm_rows:
        predil = f"1:{nrow.predilution_factor}" if nrow.predilution_factor > 1 else "—"
        stock_s = f"{nrow.stock_conc_ng_ul:.1f}" if nrow.stock_conc_ng_ul is not None else "—"
        status_label = esc(tr(f"norm.qc.{nrow.qc_status}", factor=nrow.predilution_factor))
        norm_tr.append(
            f"<tr class='qc-{esc(nrow.qc_status)}'>"
            f"<td class='chk'>&#9744;</td>"
            f"<td><b>{esc(nrow.sample)}</b></td>"
            f"<td>{esc(', '.join(nrow.wells))}</td>"
            f"<td class='num'>{stock_s}</td>"
            f"<td class='num'>{predil}</td>"
            f"<td class='num'><b>{nrow.dna_volume_ul:.2f}</b></td>"
            f"<td class='num'>{nrow.diluent_volume_ul:.2f}</td>"
            f"<td>{esc(nrow.final_summary)}</td>"
            f"<td class='qc-cell'>{status_label}</td>"
            f"</tr>"
        )

    norm_section = ""
    if norm_tr:
        norm_section = (
            f"<h2>{tr('bench.sec3')}</h2>"
            "<table class='data-table'>"
            f"<thead><tr><th>{col_done}</th><th>{esc(tr('norm.col.sample'))}</th>"
            f"<th>{esc(tr('norm.col.wells'))}</th><th>{esc(tr('norm.col.stock'))}</th>"
            f"<th>{esc(tr('norm.col.predilution'))}</th><th>{esc(tr('norm.col.dna_vol'))}</th>"
            f"<th>{esc(tr('norm.col.diluent_vol'))}</th><th>{esc(tr('norm.col.final_info'))}</th>"
            f"<th>{esc(tr('norm.col.status'))}</th></tr></thead>"
            f"<tbody>{''.join(norm_tr)}</tbody></table>"
        )

    title_str = esc(tr("bench.title", name=plate.name))
    meta_str = tr(
        "bench.meta",
        fmt=esc(plate.format.title),
        op=esc(plate.operator or "—"),
        proto=esc(recipe.name),
        dt=now_str,
    )
    sec1_str = esc(
        tr("bench.sec1", occ=len(plate.occupied_wells()), tot=plate.total_wells)
    )
    sec2_str = esc(tr("bench.sec2"))

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title_str}</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; font-size: 11px; background-color: {bg_body}; color: {fg_body}; margin: 14px; }}
  h1 {{ font-size: 16px; margin: 0 0 4px 0; color: {fg_body}; }}
  h2 {{ font-size: 13px; margin: 14px 0 6px 0; border-bottom: 1px solid {border_col}; padding-bottom: 2px; color: {fg_body}; }}
  h3 {{ font-size: 11.5px; margin: 6px 0 4px 0; color: {fg_body}; }}
  .meta {{ color: {fg_muted}; margin-bottom: 10px; }}
  table.plate-grid {{ border-collapse: collapse; width: 100%; table-layout: fixed; margin-bottom: 10px; }}
  table.plate-grid th {{ background-color: {th_grid_bg}; color: {th_grid_fg}; border: 1px solid {border_col}; padding: 2px; font-size: 10px; }}
  table.plate-grid td.well {{ border: 1px solid {border_col}; height: 34px; vertical-align: top; padding: 2px 3px; overflow: hidden; }}
  table.plate-grid td.empty {{ background-color: {empty_bg}; color: {empty_fg}; }}
  .w-coord {{ font-size: 8px; color: {coord_fg}; }}
  .w-sample {{ font-weight: 600; font-size: 9.5px; color: {fg_body}; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
  .w-assay {{ font-size: 8.5px; color: {assay_fg}; }}
  .role-tag {{ background-color: #1e293b; color: #ffffff; border-radius: 2px; padding: 0 2px; font-size: 7.5px; }}
  .mix-grid {{ display: flex; flex-wrap: wrap; gap: 12px; }}
  .mix-card {{ flex: 1 1 45%; min-width: 280px; margin-bottom: 8px; }}
  table.data-table {{ border-collapse: collapse; width: 100%; margin-bottom: 8px; }}
  table.data-table th {{ background-color: {th_data_bg}; color: {th_data_fg}; border: 1px solid {border_col}; padding: 4px 6px; text-align: left; font-size: 10px; }}
  table.data-table td {{ background-color: {td_data_bg}; color: {td_data_fg}; border: 1px solid {border_col}; padding: 4px 6px; text-align: left; }}
  table.data-table td.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
  table.data-table td.chk {{ text-align: center; width: 28px; font-size: 13px; }}
  tr.qc-ok td.qc-cell {{ background-color: {qc_ok_bg}; color: {qc_ok_fg}; }}
  tr.qc-too_concentrated td {{ background-color: {qc_warn_bg}; color: {qc_warn_fg}; }}
  tr.qc-too_dilute td {{ background-color: {qc_err_bg}; color: {qc_err_fg}; }}
  @media print {{
    body {{ margin: 5mm; }}
  }}
</style>
</head>
<body>
  <h1>{title_str}</h1>
  <div class="meta">{meta_str}</div>

  <h2>{sec1_str}</h2>
  <table class="plate-grid">
    {''.join(grid_rows_html)}
  </table>

  <h2>{sec2_str}</h2>
  <div class="mix-grid">
    {''.join(mix_blocks) if mix_blocks else f'<p>{esc(tr("mix.no_assays"))}</p>'}
  </div>

  {norm_section}
</body>
</html>
"""


def export_bench_sheet_html(
    plate: Plate,
    recipe: MixRecipe,
    norm_config: NormalizationConfig,
    path: str | Path,
) -> None:
    html_text = render_bench_sheet_html(plate, recipe, norm_config, dark=False)
    Path(path).write_text(html_text, encoding="utf-8")
